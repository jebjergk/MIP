from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt, Emu

MIP_BLUE = RGBColor(48, 105, 240)
MIP_DARK = RGBColor(15, 23, 42)
MIP_SLATE = RGBColor(51, 65, 85)
MIP_BG = RGBColor(245, 248, 252)
MIP_TEXT = RGBColor(31, 41, 55)
MIP_MUTED = RGBColor(95, 107, 122)
MIP_GREEN = RGBColor(46, 125, 50)
MIP_AMBER = RGBColor(146, 95, 0)
MIP_RED = RGBColor(198, 40, 40)
MIP_PURPLE = RGBColor(106, 27, 154)
MIP_TEAL = RGBColor(0, 121, 107)
WHITE = RGBColor(255, 255, 255)


def _set_bg(slide, color: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _draw_header(slide, title: str, subtitle: str, section: str = "MIP") -> None:
    top = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.78))
    top.fill.solid()
    top.fill.fore_color.rgb = MIP_DARK
    top.line.fill.background()

    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0.75), Inches(13.333), Inches(0.06))
    accent.fill.solid()
    accent.fill.fore_color.rgb = MIP_BLUE
    accent.line.fill.background()

    section_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.8), Inches(0.14), Inches(2.3), Inches(0.42))
    section_box.fill.solid()
    section_box.fill.fore_color.rgb = RGBColor(36, 48, 71)
    section_box.line.color.rgb = MIP_BLUE
    sec_tf = section_box.text_frame
    sec_tf.text = section
    sec_tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    sec_tf.paragraphs[0].font.size = Pt(11)
    sec_tf.paragraphs[0].font.bold = True
    sec_tf.paragraphs[0].font.color.rgb = RGBColor(225, 233, 246)

    title_box = slide.shapes.add_textbox(Inches(0.7), Inches(1.02), Inches(8.7), Inches(0.55))
    title_tf = title_box.text_frame
    title_tf.text = title
    title_tf.paragraphs[0].font.size = Pt(28)
    title_tf.paragraphs[0].font.bold = True
    title_tf.paragraphs[0].font.color.rgb = MIP_DARK

    subtitle_box = slide.shapes.add_textbox(Inches(0.7), Inches(1.58), Inches(11.8), Inches(0.38))
    subtitle_tf = subtitle_box.text_frame
    subtitle_tf.text = subtitle
    subtitle_tf.paragraphs[0].font.size = Pt(13)
    subtitle_tf.paragraphs[0].font.color.rgb = MIP_MUTED


def _add_footer(slide, text: str = "MIP | Ingestion -> Evidence -> Decision -> Trade") -> None:
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.18), Inches(13.333), Inches(0.32))
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor(238, 242, 248)
    bar.line.fill.background()
    t = slide.shapes.add_textbox(Inches(0.6), Inches(7.22), Inches(10.0), Inches(0.2))
    tf = t.text_frame
    tf.text = text
    tf.paragraphs[0].font.size = Pt(10)
    tf.paragraphs[0].font.color.rgb = RGBColor(110, 123, 141)


def _flow_box(slide, x, y, w, h, title, narrative, accent_color, title_size=10, narr_size=8):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    card.fill.solid()
    card.fill.fore_color.rgb = WHITE
    card.line.color.rgb = RGBColor(215, 220, 228)

    edge = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.06), Inches(h))
    edge.fill.solid()
    edge.fill.fore_color.rgb = accent_color
    edge.line.fill.background()

    tf = card.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    p0 = tf.paragraphs[0]
    p0.text = title
    p0.font.bold = True
    p0.font.size = Pt(title_size)
    p0.font.color.rgb = accent_color
    p0.space_after = Pt(2)

    p1 = tf.add_paragraph()
    p1.text = narrative
    p1.font.size = Pt(narr_size)
    p1.font.color.rgb = MIP_TEXT
    p1.space_before = Pt(1)


def _arrow_right(slide, x, y, w=0.18, h=0.32):
    arrow = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(x), Inches(y), Inches(w), Inches(h))
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = MIP_BLUE
    arrow.line.fill.background()


def _arrow_down(slide, x, y, w=0.32, h=0.22):
    arrow = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = RGBColor(183, 194, 214)
    arrow.line.fill.background()


def _phase_label(slide, x, y, w, text, bg_color):
    chip = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.26))
    chip.fill.solid()
    chip.fill.fore_color.rgb = bg_color
    chip.line.fill.background()
    tf = chip.text_frame
    tf.text = text
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    tf.paragraphs[0].font.size = Pt(9)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = WHITE


def build_workflow_slide(output_path: Path) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_bg(slide, MIP_BG)
    _draw_header(
        slide,
        "Symbol Lifecycle: Ingestion to Trade",
        "Complete workflow for a single symbol through the MIP pipeline, showing every gate and decision point",
        "ARCHITECTURE",
    )

    # --- ROW 1: Research Phase (y ~ 2.0 - 3.55) ---
    row1_y = 2.15
    box_h = 1.35
    box_w = 1.78
    gap = 0.08
    arrow_w = 0.16

    _phase_label(slide, 0.45, 1.92, 1.5, "RESEARCH", MIP_BLUE)

    # Box 1: Data Ingestion
    x1 = 0.45
    _flow_box(slide, x1, row1_y + 0.32, box_w, box_h,
              "1. Data Ingestion",
              "IBKR Gateway fetches daily bars (OHLCV) for all universe symbols. Merged into MARKET_BARS.",
              MIP_BLUE, title_size=9, narr_size=7.5)

    _arrow_right(slide, x1 + box_w + gap, row1_y + 0.32 + box_h / 2 - 0.16, arrow_w, 0.3)

    # Box 2: Pattern Detection
    x2 = x1 + box_w + gap + arrow_w + gap
    _flow_box(slide, x2, row1_y + 0.32, box_w, box_h,
              "2. Pattern Detection",
              "Detectors fire for each pattern type: Momentum, Mean-Reversion, Bearish Momentum.",
              MIP_BLUE, title_size=9, narr_size=7.5)

    _arrow_right(slide, x2 + box_w + gap, row1_y + 0.32 + box_h / 2 - 0.16, arrow_w, 0.3)

    # Box 3: Signal Scoring
    x3 = x2 + box_w + gap + arrow_w + gap
    _flow_box(slide, x3, row1_y + 0.32, box_w, box_h,
              "3. Signal Scoring",
              "Each signal scored and logged to RECOMMENDATION_LOG with pattern details and deviation metrics.",
              MIP_BLUE, title_size=9, narr_size=7.5)

    _arrow_right(slide, x3 + box_w + gap, row1_y + 0.32 + box_h / 2 - 0.16, arrow_w, 0.3)

    # Box 4: Outcome Evaluation
    x4 = x3 + box_w + gap + arrow_w + gap
    _flow_box(slide, x4, row1_y + 0.32, box_w, box_h,
              "4. Outcome Evaluation",
              "Realized returns measured at H1, H3, H5, H10, H20 bars. Hit rates and avg returns computed.",
              MIP_BLUE, title_size=9, narr_size=7.5)

    _arrow_right(slide, x4 + box_w + gap, row1_y + 0.32 + box_h / 2 - 0.16, arrow_w, 0.3)

    # Box 5: Trust Gating
    x5 = x4 + box_w + gap + arrow_w + gap
    _flow_box(slide, x5, row1_y + 0.32, box_w + 0.3, box_h,
              "5. Trust Gating",
              "Pattern/symbol passes if: hit rate >= threshold, avg return >= min, sample size sufficient. Trust label: TRUSTED / UNTRUSTED.",
              MIP_GREEN, title_size=9, narr_size=7.5)

    # --- Transition arrow down from row 1 to row 2 ---
    _arrow_down(slide, x5 + 0.7, row1_y + 0.32 + box_h + 0.06, 0.3, 0.22)

    # --- ROW 2: Decision + Execution Phase (y ~ 4.1 - 5.65) ---
    row2_y = 4.05
    _phase_label(slide, 0.45, 3.82, 1.5, "DECISIONS", MIP_PURPLE)

    # Box 6: Proposal Filter
    x6 = x5
    _flow_box(slide, x6, row2_y + 0.32, box_w + 0.3, box_h,
              "6. Proposal Filter",
              "Only MOMENTUM + MEAN_REV BULLISH pass. Bearish Momentum and Mean-Rev Bearish blocked from proposals.",
              MIP_GREEN, title_size=9, narr_size=7.5)

    _arrow_right(slide, x6 - gap - arrow_w, row2_y + 0.32 + box_h / 2 - 0.16, arrow_w, 0.3)

    # Box 7: Conflict Check (going right-to-left in row 2)
    x7 = x4
    _flow_box(slide, x7, row2_y + 0.32, box_w, box_h,
              "7. Conflict Check",
              "Query if Bearish Momentum or Mean-Rev Bearish fired for same symbol/TS. Warn committee.",
              MIP_AMBER, title_size=9, narr_size=7.5)

    _arrow_right(slide, x7 - gap - arrow_w, row2_y + 0.32 + box_h / 2 - 0.16, arrow_w, 0.3)

    # Box 8: Sim Committee
    x8 = x3
    _flow_box(slide, x8, row2_y + 0.32, box_w, box_h,
              "8. Sim Committee",
              "LLM evaluates: should_enter, size_factor, stop_loss_pct, target_return, hold_bars. Pattern-type-aware SL/TP guidance.",
              MIP_PURPLE, title_size=9, narr_size=7.5)

    _arrow_right(slide, x8 - gap - arrow_w, row2_y + 0.32 + box_h / 2 - 0.16, arrow_w, 0.3)

    _phase_label(slide, x2 - 0.15, 3.82, 1.5, "EXECUTION", MIP_GREEN)

    # Box 9: Live Committee
    x9 = x2
    _flow_box(slide, x9, row2_y + 0.32, box_w, box_h,
              "9. Live Committee",
              "6-agent roundtable with real-time pricing. Recalculates SL/TP based on current price. Final go/no-go.",
              MIP_PURPLE, title_size=9, narr_size=7.5)

    _arrow_right(slide, x9 - gap - arrow_w, row2_y + 0.32 + box_h / 2 - 0.16, arrow_w, 0.3)

    # Box 10: Trade Execution
    x10 = x1
    _flow_box(slide, x10, row2_y + 0.32, box_w, box_h,
              "10. Trade Execution",
              "IBKR bracket order placed: parent + TP leg + SL leg. Position tracked with SL/TP stored.",
              MIP_TEAL, title_size=9, narr_size=7.5)

    # --- Bottom monitoring strip ---
    mon_y = 6.0
    mon_box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.45), Inches(mon_y), Inches(12.45), Inches(0.95),
    )
    mon_box.fill.solid()
    mon_box.fill.fore_color.rgb = RGBColor(15, 23, 42)
    mon_box.line.color.rgb = RGBColor(36, 48, 65)

    edge = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.45), Inches(mon_y), Inches(0.06), Inches(0.95))
    edge.fill.solid()
    edge.fill.fore_color.rgb = MIP_RED
    edge.line.fill.background()

    tf = mon_box.text_frame
    tf.word_wrap = True
    p0 = tf.paragraphs[0]
    p0.text = "ONGOING: Exit Signal Monitoring"
    p0.font.bold = True
    p0.font.size = Pt(10)
    p0.font.color.rgb = RGBColor(226, 232, 240)

    narr_items = [
        "Bearish Momentum + Mean-Rev Bearish signals surfaced on Live Activity for held positions",
        "Committee can trigger early exit based on deteriorating thesis",
        "Payoff/giveback heuristics + crystallize at episode profit target",
    ]
    for item in narr_items:
        p = tf.add_paragraph()
        p.text = f"  {item}"
        p.font.size = Pt(8)
        p.font.color.rgb = RGBColor(203, 213, 225)
        p.space_before = Pt(0)

    _add_footer(slide)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)


if __name__ == "__main__":
    deck_path = Path(__file__).with_name("MIP_Symbol_Lifecycle_Workflow.pptx")
    build_workflow_slide(deck_path)
    print(f"Created: {deck_path}")
