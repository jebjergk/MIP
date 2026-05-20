#!/usr/bin/env python3
"""Build MIP introduction deck (.pptx) — palette aligned with MIP UI / cursorfiles/generate_mip_presentation.py.

Run from any cwd:
    python MIP/docs/presentations/build_mip_introduction_deck.py

Requires: pip install python-pptx

Output: MIP/docs/presentations/MIP_Introduction_Deck.pptx
Companion: mip_deck_repo_findings.md
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ── Palette (matches MIP React UI) ───────────────────────────────
NAVY_DARK = RGBColor(0x0D, 0x15, 0x40)
NAVY_LIGHT = RGBColor(0x1A, 0x28, 0x70)
ACCENT_CYAN = RGBColor(0x4F, 0xC3, 0xF7)
LINK_BLUE = RGBColor(0x0D, 0x6E, 0xFD)
TEAL = RGBColor(0x20, 0xC9, 0x97)
GREEN = RGBColor(0x19, 0x87, 0x54)
ORANGE = RGBColor(0xFD, 0x7E, 0x14)
RED = RGBColor(0xDC, 0x35, 0x45)
TEXT_DARK = RGBColor(0x1A, 0x1A, 0x2E)
TEXT_MUTED = RGBColor(0x6C, 0x75, 0x7D)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_BG = RGBColor(0xF8, 0xF9, 0xFA)
SEMI_WHITE = RGBColor(0xCC, 0xD5, 0xE0)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def add_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_shape(slide, shape_type, left, top, width, height, fill_color=None, line_rgb=None):
    s = slide.shapes.add_shape(shape_type, left, top, width, height)
    if fill_color:
        s.fill.solid()
        s.fill.fore_color.rgb = fill_color
    else:
        s.fill.background()
    if line_rgb:
        s.line.color.rgb = line_rgb
    else:
        s.line.fill.background()
    return s


def tb(
    slide,
    left,
    top,
    width,
    height,
    text,
    *,
    font_size=14,
    color=TEXT_DARK,
    bold=False,
    align=PP_ALIGN.LEFT,
    font_name="Segoe UI",
    anchor=MSO_ANCHOR.TOP,
):
    box = slide.shapes.add_textbox(left, top, width, height)
    box.text_frame.word_wrap = True
    box.text_frame.vertical_anchor = anchor
    p = box.text_frame.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(font_size)
    r.font.color.rgb = color
    r.font.bold = bold
    r.font.name = font_name


def footer_disclaimer(slide, text: str):
    tb(
        slide,
        Inches(0.6),
        Inches(6.92),
        Inches(12),
        Inches(0.45),
        text,
        font_size=11,
        color=ORANGE,
        bold=False,
        align=PP_ALIGN.LEFT,
    )


def slide_dark_header(slide, title: str):
    add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), SLIDE_W, Inches(1.08), NAVY_DARK)
    tb(slide, Inches(0.75), Inches(0.22), Inches(11.8), Inches(0.75), title, font_size=28, color=WHITE, bold=True)
    add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(0.75), Inches(0.93), Inches(1.6), Pt(3), ACCENT_CYAN)


def content_slide(title: str, subtitle: str | None, bullets: list[str], *, foot=None):
    slide = blank()
    add_bg(slide, WHITE)
    slide_dark_header(slide, title)
    y = Inches(1.22)
    if subtitle:
        tb(slide, Inches(0.75), y, Inches(11.8), Inches(0.85), subtitle, font_size=14, color=NAVY_DARK, bold=True)
        y = Inches(2.05)
    body = "\n".join("\u2022  " + b for b in bullets)
    tb(slide, Inches(0.75), y, Inches(11.8), SLIDE_H - y - Inches(0.35), body, font_size=13, color=TEXT_DARK)
    if foot:
        footer_disclaimer(slide, foot)
    return slide


prs: Presentation | None = None


def blank():
    """Add slide using current presentation blank layout."""
    assert prs is not None
    return prs.slides.add_slide(prs.slide_layouts[6])


def build():
    global prs
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    # ── Slide 1 ─────────────────────────────────────────────────
    s1 = blank()
    add_bg(s1, NAVY_DARK)
    for i in range(0, 14):
        add_shape(s1, MSO_SHAPE.RECTANGLE, Inches(i), Inches(0), Pt(0.5), SLIDE_H, RGBColor(0x14, 0x1E, 0x50))
    tb(
        s1,
        Inches(2.2),
        Inches(2.55),
        Inches(8.9),
        Inches(1.1),
        "MIP — Market Intelligence Platform",
        font_size=40,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_shape(s1, MSO_SHAPE.RECTANGLE, Inches(5.5), Inches(3.75), Inches(2.3), Pt(4), ACCENT_CYAN)
    tb(
        s1,
        Inches(2),
        Inches(3.98),
        Inches(9.3),
        Inches(1.35),
        "A private AI-assisted market research and systematic trading platform.",
        font_size=18,
        color=SEMI_WHITE,
        align=PP_ALIGN.CENTER,
    )
    tb(
        s1,
        Inches(2),
        Inches(5.55),
        Inches(9.3),
        Inches(0.8),
        "Evidence-first pipelines in Snowflake  \u2022  Bounded AI narration  \u2022  Live execution discipline",
        font_size=14,
        color=TEXT_MUTED,
        align=PP_ALIGN.CENTER,
    )

    add_shape(s1, MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.12), SLIDE_W, Inches(0.38), NAVY_LIGHT)
    add_shape(s1, MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.12), Inches(3.8), Pt(4), ACCENT_CYAN)

    content_slide(
        "Why MIP Exists",
        "Operators drown in signal noise; systems need repeatable proof and oversight.",
        [
            "Charts and scans produce disconnected ideas — scaling coverage without drowning in contradiction is hard.",
            "Pure rule bots are brittle: markets shift; rules need statistically visible aging and diagnostics.",
            "Generic AI trading chat can sound confident without evidence trails — that is risky where capital is involved.",
            "Live trading needs inspectable decisions: what data, what policy, what broker truth — with audit lineage.",
        ],
    )

    content_slide(
        "The MIP Philosophy",
        "Deterministic substrate first — AI explains, challenges, narrates — it does not replace truth.",
        [
            "Deterministic evidence first: bars, setups, validated outcomes in Snowflake APP/MART schemas.",
            "AI as a reasoning layer atop structured facts — not undocumented magic.",
            "Snowflake as analytical truth layer: procedures, tables, mart views compose the playbook.",
            "IBKR (when ADAPTER_MODE allows) as execution-layer truth aligned with LIVE_PORTFOLIO_CONFIG.",
            "Every decision pathway should trace to inspectable rows, reason codes, and operator controls.",
        ],
    )

    # Slide 4 — lifecycle diagram (text ladder + boxes)
    s4 = blank()
    add_bg(s4, WHITE)
    slide_dark_header(s4, "MIP Lifecycle Overview")
    tb(
        s4,
        Inches(0.75),
        Inches(1.2),
        Inches(11.5),
        Inches(0.95),
        "One chain from market bars to disciplined live operations.",
        font_size=14,
        color=NAVY_DARK,
        bold=True,
    )
    stages = [
        ("Ingest", "Bars \u2192 MARKET_BARS", LINK_BLUE),
        ("Returns / legacy signal", "MARKET_RETURNS, momentum RECs", TEAL),
        ("Structural discovery", "Levels, setups, regimes, trust", GREEN),
        ("Proposals", "Deterministic + Phase 4 board", ACCENT_CYAN),
        ("Committee / LPA", "Hearings, LPA cockpit", ORANGE),
        ("Broker", "Orders, fills, reconcile", NAVY_LIGHT),
    ]
    bx, by, bw, bh = Inches(0.45), Inches(2.35), Inches(1.85), Inches(1.1)
    for i, (label, sub, clr) in enumerate(stages):
        x = bx + i * Inches(1.93)
        add_shape(s4, MSO_SHAPE.ROUNDED_RECTANGLE, x, by, bw, bh, clr)
        tb(s4, x + Inches(0.08), by + Inches(0.1), bw - Inches(0.16), Inches(0.42), label, font_size=12, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
        tb(s4, x + Inches(0.06), by + Inches(0.52), bw - Inches(0.12), Inches(0.52), sub, font_size=9, color=RGBColor(0xEE, 0xEE, 0xFF), align=PP_ALIGN.CENTER)
        if i < len(stages) - 1:
            add_shape(
                s4,
                MSO_SHAPE.RIGHT_ARROW,
                x + bw + Inches(0.06),
                by + Inches(0.4),
                Inches(0.35),
                Inches(0.28),
                ACCENT_CYAN,
            )

    tb(
        s4,
        Inches(0.65),
        Inches(3.82),
        Inches(11.8),
        Inches(3.05),
        "Closer loop:\n\u2022  Daily TASK_RUN_DAILY_PIPELINE \u2192 SP_RUN_DAILY_PIPELINE orchestrates ingestion, returns,"
        "\n    momentum recommendation generation/outcomes, portfolio simulation,\n"
        "    morning brief persistence,\n\u2022  then SP_RUN_STRUCTURAL_DAILY_PIPELINE (levels → setups → evaluate → trust → propose).\n"
        "\u2022  Phase 4 run_board stages Cortex agents publishing STRUCTURAL_TRADE_PROPOSALS (BOARD_RUN_ID set).\n"
        "\u2022  Parallel Worlds + daily position verdict (SP_RUN_DAILY_POSITION_VERDICT when enabled) tighten governance.",
        font_size=12,
        color=TEXT_DARK,
    )

    content_slide(
        "Data Ingestion & Daily Pipeline",
        None,
        [
            "TASK_RUN_DAILY_PIPELINE (Snowflake task) invokes SP_RUN_DAILY_PIPELINE after US equity close window (cron in 150_task_run_daily_training.sql).",
            "SP_PIPELINE_INGEST \u2192 SP_INGEST_MARKET_BARS: routes by MARKET_DATA_PROVIDER_DEFAULT — IB ingest via cursorfiles/ingest_ibkr_bars.py is primary path in ops docs; Alpha Vantage proc remains available.",
            "MARKET_BARS refreshed in MART; SP_PIPELINE_REFRESH_RETURNS maintains MARKET_RETURNS.",
            "Momentum path: SP_GENERATE_MOMENTUM_RECS \u2192 RECOMMENDATION_LOG; evaluate \u2192 RECOMMENDATION_OUTCOMES (H1,H3,H5,H10,H20).",
            "Portfolio simulation SP_RUN_PORTFOLIO_SIMULATION updates PORTFOLIO_* tables.",
            "Morning digest path: SP_WRITE_MORNING_BRIEF merges V_MORNING_BRIEF_JSON into AGENT_OUT.MORNING_BRIEF.",
            "Manual operator path: mip_ui_api POST /ib/daily-job/run (+ optional synth_intraday_daily, pipeline, run_board flags).",
        ],
    )

    content_slide(
        "Training & Evidence Accumulation",
        None,
        [
            "Legacy momentum trainers: horizons and HIT coverage in RECOMMENDATION_OUTCOMES; trust signals via mart/training KPI views.",
            "Structural track: STRUCTURAL_SETUP_EVENTS linked to STRUCTURAL_SETUP_OUTCOMES once forward bars exist (SP_EVALUATE_STRUCTURAL_OUTCOMES).",
            "STRUCTURAL_SETUP_TRUST rolls up MEANINGFUL_HIT_RATE, path survival, MFE/MAE style metrics per family / market."
            "\n(TRUST_LABEL bands such as TRUSTED / PROVISIONAL / RESEARCH depending on thresholds in SQL.)",
            "Meaningful move thresholds encode what counts as structural success versus noise — parameterized in APP strategy tables/views.",
            "Operator surfaces: /structural-training, /training (digest + maturity vocabulary). Ask MIP can explain rows in plain language.",
        ],
    )

    content_slide(
        "Structural Pattern Detection",
        None,
        [
            "Daily structural DAG (SP_RUN_STRUCTURAL_DAILY_PIPELINE) detects pivot levels SP_DETECT_STRUCTURAL_LEVELS,"
            "\ncomputes structural state/regime compatibility, evaluates structural setups.",
            "Setup families surfaced in mart views include breakout/retest, wick rejections at support/resistance,"
            "\ntrend pullbacks, failed breakout reversals — see SETUP_NARRATIVE cases in V_STRUCTURAL_TIMELINE_SETUPS.",
            "Outputs track LEVEL_PRICE, ENTRY_ZONE_LOW/HIGH, invalidation context, wick and multi-bar confirmation scores.",
            "Market structure map view (v_symbol_market_structure_map) adds BOS/CHOCH/wick-probe style tags for richer context.",
        ],
    )

    content_slide(
        "Structural Timeline",
        "Market memory for a symbol: recent structure, levels, events, proposal linkage.",
        [
            "UI route /structural-timeline \u2192 FastAPI /structural-timeline/detail?setup_event_id=…",
            "Mart rail V_STRUCTURAL_TIMELINE_EVENTS unions state changes, setup lifecycle, proposal_create, new levels.",
            "V_STRUCTURAL_TIMELINE_LEVELS reads STRUCTURAL_LEVEL_CACHE for scored support/resistance bands.",
            "Setups view joins trust policy, proposal counts, regime tags, narrative text for education.",
        ],
    )

    content_slide(
        "Agentic Proposal Board",
        "Evidence pack \u2192 specialist agents \u2192 chair \u2192 durable proposals (separate from deterministic SP_PROPOSE_STRUCTURAL_TRADES).",
        [
            "Phase 4 entry: MIP/scripts/proposal_board_phase4/run_board.py stages dossiers (V_PROPOSAL_BOARD_*), parallel specialists, chair, publication policy.",
            "Published rows appear in STRUCTURAL_TRADE_PROPOSALS with BOARD_RUN_ID; timeline/agentic narratives key off that discriminator.",
            "Operator trigger commonly chains after pipeline via POST /ib/daily-job/run (run_proposal_board flag).",
            "Short-side publication guarded by LIVE_PORTFOLIO_CONFIG.IBKR_ACCOUNT_MODE (\u2260 UNKNOWN / fail-closed logic in orchestrator).",
        ],
    )

    content_slide(
        "Committee & Agentic Revalidation",
        "Re-open a thesis: still valid, degraded, or closed / blocked by policy?",
        [
            "/structural-committee + committee.py hearings: OPEN / REFRESH / COMMIT flows persist structured stance and reason scaffolding.",
            "Shadow board helpers for exploratory runs (committee_shadow_board_* endpoints).",
            "LPA inline committee exhibits embed context while acting on LIVE_ACTION rows (LpaCommittee2Exhibits).",
            "Deterministic counterpart: SP_RUN_DAILY_POSITION_VERDICT + Position Health (/position-health) scores thesis_integrity,"
            "\nfragility vs invalidation cushions — distinct from conversational agents but feeds the same operator story.",
            "Ask: is price still near the thesis zone? Is structure repairing or snapping? Answers combine SQL verdicts + hearings + tape.",
        ],
    )

    content_slide(
        "Live Portfolio Activity (LPA)",
        "Operator cockpit for open risk, broker sync, proposals, guarded execution clicks.",
        [
            "/live-portfolio-activity surfaces pending LIVE_ACTION payloads, snapshots, stale revalidation cues.",
            "Reason-code map translates machine codes (e.g. EXECUTION_CLICK_REVALIDATION_STALE, PRICE_GUARD_FAIL, drift codes) into operator prose.",
            "Links outward to hearings, portfolios, audits when context panels mount.",
            "Keeps deterministic guardrails visible before anything hits IBKR — single pane to \"decide responsibly.\"",
        ],
    )

    content_slide(
        "Trade Configuration, Brackets & IBKR",
        "Configure intent and broker paths — mind the overloaded \"paper\" wording.",
        [
            "/live-portfolio-config edits LIVE_PORTFOLIO_CONFIG knobs (adapter mode, IB account ids, sizing/risk placeholders per schema).",
            "live.py composes brackets, realism checks referencing broker snapshots — separate env LIVE_EXECUTION_MODE can force IBKR or placeholder semantics.",
            "ADAPTER_MODE 'LIVE' means broker submit wiring — NOT automatically \"real-money\"; IBKR_ACCOUNT_MODE (PAPER/REAL/UNKNOWN) is explicit safety.",
            "Ingest bars + executions reconcile through IBKR-aligned scripts and audit tables referenced in LIVE router smoke paths.",
        ],
    )

    content_slide(
        "Ask MIP",
        None,
        [
            "Embedded assistant (AskMipPanel) posts conversation + route context to POST /ask/v3.",
            "Backend binds User Guide segments, glossary, optional web fallback governed by orchestrator policies (see docs/ask_mip/ask_mip_v2_design.md).",
            "Uses Snowflake Cortex COMPLETE for natural language while blocking undocumented internal claims.",
            "Operators get \"why is this badge red?\", \"explain this KPI\", onboarding without paging an engineer.",
        ],
    )

    content_slide(
        "Parallel Worlds",
        "Simulation intelligence — contrasts alternate policy scenarios vs realized path.",
        [
            "/parallel-worlds aggregates portfolio-level scenario grids, regret heatmaps, equity curves, tuning surfaces.",
            "API reads mart views such as V_PARALLEL_WORLD_REGRET joined to active PARALLEL_WORLD_SCENARIO rows.",
            "Use case: understand opportunity cost / policy sensitivity — not a substitute for broker execution truth.",
        ],
    )

    # Slide 15 — RAG design (disclaimer)
    s15 = blank()
    add_bg(s15, WHITE)
    slide_dark_header(s15, "Literature-Informed Intelligence (Planned Layer)")
    tb(
        s15,
        Inches(0.75),
        Inches(1.18),
        Inches(11.6),
        Inches(1.25),
        "Turn external trading literature into a controlled knowledge layer — not a trade-decision authority.",
        font_size=14,
        color=NAVY_DARK,
        bold=True,
    )
    rag_bullets = [
        "Ingest books / research → extract setup concepts, context rules, invalidation templates, failure modes, trade-management heuristics.",
        "Normalize into structured concept cards with provenance and confidence metadata.",
        "Tag each card: observable with existing MIP bars/structure versus narrative-only.",
        "Inject cards into Phase 4 / committee / revalidation dialogs as debate material — challenge weak thesis, articulate failure symmetry.",
        "Binding decisions remain deterministic evidence + broker truth + risk policies + STRUCTURAL_SETUP_TRUST / outcomes.",
        "Illustrative style note: contextual trading literature (e.g. bar-reading frameworks) mapped to setup families (STRUCTURAL_SETUP_EVENTS.SETUP_FAMILY) conceptually.",
    ]
    tb(
        s15,
        Inches(0.75),
        Inches(2.35),
        Inches(11.6),
        Inches(4.45),
        "\n".join("\u2022  " + x for x in rag_bullets),
        font_size=12,
        color=TEXT_DARK,
    )
    footer_disclaimer(s15, "Architecture target \u2014 not yet implemented in Snowflake/repo. Companion doc logs status as Planned.")

    # Slide 16 — Cortex Code
    s16 = blank()
    add_bg(s16, WHITE)
    slide_dark_header(s16, "Engineering Accelerator: Cortex Code + Cursor")
    tb(
        s16,
        Inches(0.75),
        Inches(1.18),
        Inches(11.6),
        Inches(1.1),
        "Separate from runtime Cortex agents: tooling velocity for APP/MART SQL stewardship.",
        font_size=14,
        color=NAVY_DARK,
        bold=True,
    )
    cc_bullets = [
        "Cursor keeps repo grounding (ADRs, smoke SQL, conventions for SP_/V_/TASK_ naming, migration hygiene).",
        "Cortex Code (Snowflake IDE assistant) proposes/refactors SQL & JavaScript procs with warehouse-native idioms (VARIANT, QUALIFY, EXECUTE AS CALLER).",
        "Generate candidate smoke fragments mirroring MIP/SQL/smoke patterns; humans paste + adjust before deploy.",
        "Auto-annotate semantic intent bridging APP tables vs MART views for onboarding docs.",
        "Pair on incident playbooks: translate live.py guard findings into targeted diagnostic SELECTs.",
        "Workflow: prompt \u2192 diff review in Cursor \u2192 PR \u2192 deploy via existing MIP_ADMIN_ROLE procedures — no auto-deploy from model output.",
    ]
    tb(
        s16,
        Inches(0.75),
        Inches(2.2),
        Inches(11.6),
        Inches(4.55),
        "\n".join("\u2022  " + x for x in cc_bullets),
        font_size=12,
        color=TEXT_DARK,
    )
    footer_disclaimer(s16, "Vision practice \u2014 Snowflake Cortex Code integration not wired in this repository today.")

    content_slide(
        "What Makes MIP Different",
        None,
        [
            "Versus charting SaaS: MIP couples signals to multi-horizon outcome tables and operator-grade audit/logging (MIP_AUDIT_LOG pipeline steps).",
            "Versus brittle bots: separates evolving structural detectors from trust overlays (STRUCTURAL_SETUP_TRUST) plus committee/LPA veto paths.",
            "Versus ChatGPT trader: deterministic rows + Cortex narration + explicit reason codes guard against confident hallucinations.",
            "Versus opaque quant boxes: dossier/board pipeline documents evidence packs feeding chair decisions; Parallel Worlds exposes scenario regret.",
            "Hybrid: deterministic core + cautious agent choreography + human cockpit for final intent.",
        ],
    )

    # Slide 18 — audience columns
    s18 = blank()
    add_bg(s18, WHITE)
    slide_dark_header(s18, "Audience-Specific Takeaways")
    col_specs = [
        (
            "Traders",
            GREEN,
            (
                "Structural timeline + trust overlays before sizing.\n"
                "Committee/LPA revalidation when tape drifts.\n"
                "Parallel Worlds for policy regret — not execution truth."
            ),
        ),
        (
            "AI / ML curious",
            LINK_BLUE,
            (
                "Hybrid: Snowflake evidence + Cortex agents + human cockpit.\n"
                "BOARD_RUN_ID marks agentic proposals vs SQL-only path.\n"
                "Planned literature cards = debate material, not veto authority."
            ),
        ),
        (
            "Finance / risk leaders",
            TEAL,
            (
                "Governance: schemas, audit log, explicit IBKR_ACCOUNT_MODE.\n"
                "Traceability from bar row \u2192 setup \u2192 proposal \u2192 order intent.\n"
                "Automation bounded by reason codes and broker reconciliation."
            ),
        ),
        (
            "Technical explorers",
            NAVY_LIGHT,
            (
                "Read SP_RUN_DAILY_PIPELINE + SP_RUN_STRUCTURAL_DAILY_PIPELINE.\n"
                "FastAPI read models + React cockpit mirror mart truth.\n"
                "Smoke SQL under MIP/SQL/smoke reproduces checkpoints."
            ),
        ),
    ]
    cw = Inches(2.92)
    gap = Inches(0.2)
    x0 = Inches(0.45)
    y0 = Inches(1.35)
    ch = Inches(5.05)
    for idx, (title, color, blob) in enumerate(col_specs):
        x = x0 + idx * (cw + gap)
        add_shape(s18, MSO_SHAPE.ROUNDED_RECTANGLE, x, y0, cw, Inches(0.38), color)
        tb(s18, x, y0 + Inches(0.05), cw, Inches(0.3), title, font_size=12, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
        card = add_shape(s18, MSO_SHAPE.ROUNDED_RECTANGLE, x, y0 + Inches(0.43), cw, ch - Inches(0.43), LIGHT_BG)
        card.line.color.rgb = color
        card.line.width = Pt(1.5)
        tb(s18, x + Inches(0.14), y0 + Inches(0.55), cw - Inches(0.28), ch - Inches(0.7), blob, font_size=11, color=TEXT_DARK)

    # Slide 19 — roadmap pillars
    s19 = blank()
    add_bg(s19, WHITE)
    slide_dark_header(s19, "Current State vs Roadmap")
    pillars = [
        (
            "Implemented now",
            GREEN,
            "SP_RUN_DAILY_PIPELINE + structural DAG\nMomentum + structural trust tables\nPhase 4 run_board\nLPA + live.py + LIVE_PORTFOLIO_CONFIG\nParallel Worlds mart + UI\nAsk MIP /ask/v3 + Cortex COMPLETE",
        ),
        (
            "Partial / wiring",
            ORANGE,
            "Intraday TASK_RUN_INTRADAY_PIPELINE suspended\nCommittee2 feature depth env-specific\nAdapter vs IBKR_ACCOUNT_MODE wording cleanup backlog\nOptional PW panels tolerant of sparse data",
        ),
        (
            "Future vision",
            LINK_BLUE,
            "Literature RAG concept-card layer (Slide 15)\nCortex Code + Cursor SQL velocity loop\nDeeper retrieval inside Ask beyond static guide",
        ),
    ]
    for i, (label, clr, txt) in enumerate(pillars):
        x = Inches(0.55 + i * 4.08)
        add_shape(s19, MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.35), Inches(3.85), Inches(0.4), clr)
        tb(s19, x, Inches(1.4), Inches(3.85), Inches(0.33), label, font_size=13, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
        tb(s19, x + Inches(0.15), Inches(1.88), Inches(3.55), Inches(4.95), txt, font_size=12, color=TEXT_DARK)

    # Slide 20 — closing
    s20 = blank()
    add_bg(s20, NAVY_DARK)
    tb(
        s20,
        Inches(1.2),
        Inches(2.45),
        Inches(10.9),
        Inches(1.1),
        "Closing Vision",
        font_size=36,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_shape(s20, MSO_SHAPE.RECTANGLE, Inches(5.0), Inches(3.6), Inches(3.3), Pt(4), ACCENT_CYAN)
    tb(
        s20,
        Inches(1.25),
        Inches(4.05),
        Inches(10.85),
        Inches(2.4),
        "MIP is an evolving AI-assisted trading operating system \u2014 research, deterministic evidence,"
        "\ncommittee-grade challenge, guarded live execution oversight, simulation feedback,"
        "\nand disciplined learning loops anchored in Snowflake + broker truth.",
        font_size=18,
        color=SEMI_WHITE,
        align=PP_ALIGN.CENTER,
    )
    tb(
        s20,
        Inches(1.25),
        Inches(6.55),
        Inches(10.85),
        Inches(0.7),
        "Market Intelligence Platform  \u2022  May 2026  \u2022  See mip_deck_repo_findings.md for object-level traceability.",
        font_size=12,
        color=TEXT_MUTED,
        align=PP_ALIGN.CENTER,
    )
    add_shape(s20, MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.12), SLIDE_W, Inches(0.38), NAVY_LIGHT)
    add_shape(s20, MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.12), Inches(3.8), Pt(4), ACCENT_CYAN)

    return prs


def main() -> None:
    global prs
    out_dir = Path(__file__).resolve().parent
    built = build()
    out_path = out_dir / "MIP_Introduction_Deck.pptx"
    built.save(out_path.as_posix())
    print(f"Wrote {out_path}")


if __name__ == "__main__":
    main()